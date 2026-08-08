from collections import OrderedDict
from datetime import datetime, timedelta, date
from decimal import Decimal, InvalidOperation
from functools import wraps
from io import BytesIO
from flask import render_template, redirect, url_for, flash, request, abort, jsonify, Response
from flask_login import current_user, logout_user
from sqlalchemy.orm import joinedload, subqueryload
from app.main import main_bp
from app.main.forms import PatientForm
from app.extensions import db
from app.models import (Patient, Examination, ExaminationAnalysis,
                        ExaminationDirection, ExaminationAnalysisTool, ExaminationBlank,
                        DoctorDirection, DoctorDirectionCategory,
                        Analysis, AnalysisTool, Blank,
                        AnalysisToolSubcategory,
                        CombinedAnalysis, User, EarningPlan, user_directions)


def role_required(*roles):
    """Decorator factory for main-section routes.
    If roles is empty — allows any non-admin authenticated user.
    If roles are specified — user must have one of them.
    """
    def decorator(f):
        @wraps(f)
        def decorated(*args, **kwargs):
            if not current_user.is_authenticated:
                return redirect(url_for('auth.login'))
            if current_user.is_administrator():
                return redirect(url_for('admin.dashboard'))
            if not current_user.is_active:
                logout_user()
                flash('Siziň ulanyjyňyz bloklanan.', 'danger')
                return redirect(url_for('auth.login'))
            if roles and current_user.role not in roles:
                flash('Siz bu bölege girip bilmeýärsiňiz.', 'danger')
                return redirect(url_for('main.dashboard'))
            return f(*args, **kwargs)
        return decorated
    return decorator


CASHIER_ROLES = ('cashier', 'senior_cashier')

main_required = role_required()
patients_required = role_required('registrar', 'doctor')
examinations_view_required = role_required('registrar', 'doctor', 'cashier', 'analysis_responsible', 'senior_cashier')
reports_required = role_required('cashier', 'senior_cashier')


# ── Dashboard ─────────────────────────────────────────────────────────────────

@main_bp.route('/')
@main_bp.route('/dashboard')
@main_required
def dashboard():
    stats = None
    if current_user.role in ('registrar', 'doctor'):
        stats = {
            'total': Patient.query.count(),
            'active': Patient.query.filter_by(is_active=True).count(),
            'blocked': Patient.query.filter_by(is_active=False).count(),
        }
    return render_template('main/dashboard.html', stats=stats)


# ── Patients list ─────────────────────────────────────────────────────────────

@main_bp.route('/patients')
@patients_required
def patients_list():
    search = request.args.get('q', '').strip()
    status_filter = request.args.get('status', '').strip()

    query = Patient.query

    if search:
        like = f'%{search}%'
        query = query.filter(
            db.or_(
                Patient.full_name.ilike(like),
                Patient.passport_number.ilike(like),
                Patient.insurance_number.ilike(like),
                Patient.citizenship.ilike(like),
            )
        )

    if status_filter == 'active':
        query = query.filter_by(is_active=True)
    elif status_filter == 'blocked':
        query = query.filter_by(is_active=False)

    page = request.args.get('page', 1, type=int)
    pagination = query.order_by(Patient.full_name).paginate(page=page, per_page=15, error_out=False)

    return render_template(
        'main/patients/list.html',
        patients=pagination.items,
        pagination=pagination,
        search=search,
        status_filter=status_filter,
    )


# ── Create patient ────────────────────────────────────────────────────────────

@main_bp.route('/patients/create', methods=['GET', 'POST'])
@patients_required
def patients_create():
    form = PatientForm()
    if form.validate_on_submit():
        patient = Patient(
            full_name=form.full_name.data.strip(),
            birth_year=form.birth_year.data,
            citizenship=form.citizenship.data.strip(),
            home_address=form.home_address.data.strip(),
            passport_number=form.passport_number.data.strip() or None,
            insurance_number=form.insurance_number.data.strip() or None,
        )
        db.session.add(patient)
        db.session.commit()
        flash(f'Syrkaw «{patient.full_name}» döredilen.', 'success')
        return redirect(url_for('main.patients_list'))

    return render_template('main/patients/create.html', form=form)


# ── Edit patient ──────────────────────────────────────────────────────────────

@main_bp.route('/patients/<int:patient_id>/edit', methods=['GET', 'POST'])
@patients_required
def patients_edit(patient_id):
    patient = db.session.get(Patient, patient_id)
    if patient is None:
        abort(404)

    form = PatientForm(obj=patient, editing_patient=patient)

    if form.validate_on_submit():
        patient.full_name = form.full_name.data.strip()
        patient.birth_year = form.birth_year.data
        patient.citizenship = form.citizenship.data.strip()
        patient.home_address = form.home_address.data.strip()
        patient.passport_number = form.passport_number.data.strip() or None
        patient.insurance_number = form.insurance_number.data.strip() or None
        db.session.commit()
        flash(f'Syrkaw «{patient.full_name}» maglumatlary täzelenen.', 'success')
        return redirect(url_for('main.patients_list'))

    return render_template('main/patients/edit.html', form=form, patient=patient)


# ── Toggle block/unblock patient ──────────────────────────────────────────────

@main_bp.route('/patients/<int:patient_id>/toggle', methods=['POST'])
@patients_required
def patients_toggle(patient_id):
    patient = db.session.get(Patient, patient_id)
    if patient is None:
        abort(404)

    patient.is_active = not patient.is_active
    db.session.commit()

    action = 'aktiw' if patient.is_active else 'bloklanan'
    flash(f'Syrkaw «{patient.full_name}» {action}.', 'success')
    return redirect(url_for('main.patients_list'))


# ── Patient examination history ───────────────────────────────────────────────

@main_bp.route('/patients/<int:patient_id>/history')
@patients_required
def patients_history(patient_id):
    patient = db.session.get(Patient, patient_id)
    if patient is None:
        abort(404)
    examinations = (
        Examination.query
        .filter_by(patient_id=patient_id)
        .options(
            subqueryload(Examination.exam_analyses).joinedload(ExaminationAnalysis.analysis),
            subqueryload(Examination.exam_directions).joinedload(ExaminationDirection.direction),
            subqueryload(Examination.exam_directions).joinedload(ExaminationDirection.doctor),
            subqueryload(Examination.exam_tools).joinedload(ExaminationAnalysisTool.tool),
            subqueryload(Examination.exam_blanks).joinedload(ExaminationBlank.blank),
            joinedload(Examination.created_by),
            joinedload(Examination.paid_by),
        )
        .order_by(Examination.created_at.desc())
        .all()
    )
    return render_template('main/patients/history.html',
                           patient=patient, examinations=examinations)


# ── Patients search API ───────────────────────────────────────────────────────

@main_bp.route('/patients/search')
@patients_required
def patients_search():
    q = request.args.get('q', '').strip()
    if not q or len(q) < 2:
        return jsonify([])
    patients = (Patient.query
                .filter(Patient.is_active == True,
                        Patient.full_name.ilike(f'%{q}%'))
                .order_by(Patient.full_name)
                .limit(20)
                .all())
    return jsonify([{
        'id': p.id,
        'full_name': p.full_name,
        'birth_year': p.birth_year,
        'passport_number': p.passport_number or '',
        'insurance_number': p.insurance_number or '',
    } for p in patients])


# ── Examinations helpers ──────────────────────────────────────────────────────

def _exam_form_context():
    """Return data needed to render create/edit examination form."""
    all_tools = (AnalysisTool.query
                 .filter_by(is_active=True)
                 .order_by(AnalysisTool.name)
                 .all())
    analysis_tools_map = {}
    for t in all_tools:
        for a in t.analyses:
            analysis_tools_map.setdefault(a.id, []).append(t.id)

    all_blanks = (Blank.query
                  .filter_by(is_active=True)
                  .order_by(Blank.name)
                  .all())
    analysis_blanks_map = {}
    for b in all_blanks:
        for a in b.analyses:
            analysis_blanks_map.setdefault(a.id, []).append(b.id)

    directions = (DoctorDirection.query
                  .filter_by(is_active=True)
                  .order_by(DoctorDirection.name)
                  .all())
    analysis_directions_map = {}
    for d in directions:
        for a in d.analyses:
            analysis_directions_map.setdefault(a.id, []).append(d.id)

    return {
        'analyses': Analysis.query.filter_by(is_active=True).order_by(Analysis.name).all(),
        'combined_analyses': CombinedAnalysis.query.filter_by(is_active=True).options(joinedload(CombinedAnalysis.analyses)).order_by(CombinedAnalysis.name).all(),
        'directions': directions,
        'doctors': User.query.filter(
            User.role.in_(['doctor', 'analysis_responsible']),
            User.is_active == True,
        ).options(joinedload(User.directions)).order_by(User.full_name).all(),
        'all_tools': all_tools,
        'analysis_tools_map': analysis_tools_map,
        'all_blanks': all_blanks,
        'analysis_blanks_map': analysis_blanks_map,
        'analysis_directions_map': analysis_directions_map,
    }


def _parse_exam_form():
    """Parse and validate POST data for examination form. Returns (data_dict, errors)."""
    patient_id = request.form.get('patient_id', type=int)
    analysis_ids = list(dict.fromkeys(request.form.getlist('analysis_ids', type=int)))
    direction_ids = list(dict.fromkeys(request.form.getlist('direction_ids', type=int)))
    tool_ids = list(dict.fromkeys(request.form.getlist('tool_ids', type=int)))
    blank_ids = list(dict.fromkeys(request.form.getlist('blank_ids', type=int)))

    errors = []

    if not patient_id:
        errors.append('Syrkawy saýlaň.')
    else:
        p = db.session.get(Patient, patient_id)
        if not p or not p.is_active:
            errors.append('Saýlanan syrkaw tapylmady ýa-da bloklanan.')

    if not analysis_ids and not direction_ids:
        errors.append('Iň bolmanda 1 analiz ýa-da ugur kesgitläň.')

    analysis_qtys = {aid: max(1, request.form.get(f'analysis_qty_{aid}', 1, type=int))
                     for aid in analysis_ids}
    tool_qtys = {tid: max(1, request.form.get(f'tool_qty_{tid}', 1, type=int))
                 for tid in tool_ids}
    blank_qtys = {bid: max(1, request.form.get(f'blank_qty_{bid}', 1, type=int))
                  for bid in blank_ids}

    doctor_for = {}
    if direction_ids:
        directions_by_id = {d.id: d for d in DoctorDirection.query.filter(DoctorDirection.id.in_(direction_ids)).all()}
    else:
        directions_by_id = {}
    for did in direction_ids:
        doc_id = request.form.get(f'doctor_for_{did}', type=int)
        if not doc_id:
            dir_obj = directions_by_id.get(did)
            dir_name = dir_obj.name if dir_obj else f'#{did}'
            errors.append(f'Ugur «{dir_name}» üçin lukman bellenmedik.')
        else:
            doctor_for[did] = doc_id

    data = {
        'patient_id': patient_id,
        'analysis_ids': analysis_ids,
        'analysis_qtys': analysis_qtys,
        'direction_ids': direction_ids,
        'tool_ids': tool_ids,
        'tool_qtys': tool_qtys,
        'blank_ids': blank_ids,
        'blank_qtys': blank_qtys,
        'doctor_for': doctor_for,
        'selected_patient_id': patient_id,
        'selected_analysis_qtys': analysis_qtys,
        'selected_direction_ids': set(direction_ids),
        'selected_tool_qtys': tool_qtys,
        'selected_blank_qtys': blank_qtys,
    }
    return data, errors


# ── Examinations list ─────────────────────────────────────────────────────────

@main_bp.route('/examinations')
@examinations_view_required
def examinations_list():
    search = request.args.get('q', '').strip()
    status_filter = request.args.get('status', '').strip()
    paid_filter = request.args.get('paid', '').strip()

    query = Examination.query.join(Patient, Examination.patient_id == Patient.id)

    if current_user.role == 'doctor':
        query = (query
                 .join(ExaminationDirection, ExaminationDirection.examination_id == Examination.id)
                 .filter(ExaminationDirection.doctor_id == current_user.id)
                 .distinct())
    elif current_user.role == 'analysis_responsible':
        query = (query
                 .join(ExaminationAnalysis, ExaminationAnalysis.examination_id == Examination.id)
                 .join(Analysis, Analysis.id == ExaminationAnalysis.analysis_id)
                 .filter(Analysis.responsible_id == current_user.id)
                 .distinct())

    if search:
        like = f'%{search}%'
        query = query.filter(Patient.full_name.ilike(like))

    if status_filter == 'open':
        query = query.filter(Examination.status == Examination.STATUS_OPEN)
    elif status_filter == 'closed':
        query = query.filter(Examination.status == Examination.STATUS_CLOSED)

    if paid_filter == 'paid':
        query = query.filter(Examination.is_paid == True)
    elif paid_filter == 'unpaid':
        query = query.filter(Examination.is_paid == False)

    page = request.args.get('page', 1, type=int)
    pagination = (query
        .options(
            joinedload(Examination.patient),
            joinedload(Examination.created_by),
            subqueryload(Examination.exam_analyses),
            subqueryload(Examination.exam_directions),
        )
        .order_by(Examination.created_at.desc())
        .paginate(page=page, per_page=15, error_out=False))

    return render_template(
        'main/examinations/list.html',
        examinations=pagination.items,
        pagination=pagination,
        search=search,
        status_filter=status_filter,
        paid_filter=paid_filter,
    )


# ── Create examination ────────────────────────────────────────────────────────

@main_bp.route('/examinations/create', methods=['GET', 'POST'])
@patients_required
def examinations_create():
    ctx = _exam_form_context()

    if request.method == 'POST':
        data, errors = _parse_exam_form()
        if errors:
            for e in errors:
                flash(e, 'danger')
            selected_patient = db.session.get(Patient, data['patient_id']) if data['patient_id'] else None
            return render_template('main/examinations/create.html', **ctx, **data,
                                   selected_patient=selected_patient)

        patient = db.session.get(Patient, data['patient_id'])
        exam = Examination(
            patient_id=data['patient_id'],
            created_by_id=current_user.id,
            patient_has_insurance=bool(patient.insurance_number),
        )
        db.session.add(exam)
        db.session.flush()

        analyses_by_id = {a.id: a for a in Analysis.query.filter(
            Analysis.id.in_(data['analysis_ids'])).all()} if data['analysis_ids'] else {}
        for aid in data['analysis_ids']:
            a = analyses_by_id[aid]
            db.session.add(ExaminationAnalysis(
                examination_id=exam.id,
                analysis_id=aid,
                quantity=data['analysis_qtys'].get(aid, 1),
                price=a.price,
                is_insurance=a.is_insurance,
            ))
        directions_by_id = {d.id: d for d in DoctorDirection.query.filter(
            DoctorDirection.id.in_(data['direction_ids'])).all()} if data['direction_ids'] else {}
        for did in data['direction_ids']:
            d = directions_by_id[did]
            db.session.add(ExaminationDirection(
                examination_id=exam.id,
                direction_id=did,
                doctor_id=data['doctor_for'][did],
                price=d.price,
                is_insurance=d.is_insurance,
            ))
        tools_by_id = {t.id: t for t in AnalysisTool.query.filter(
            AnalysisTool.id.in_(data['tool_ids'])).all()} if data['tool_ids'] else {}
        for tid in data['tool_ids']:
            t = tools_by_id[tid]
            db.session.add(ExaminationAnalysisTool(
                examination_id=exam.id,
                tool_id=tid,
                quantity=data['tool_qtys'].get(tid, 1),
                price=t.total_price,
                is_insurance=t.is_insurance,
            ))
        blanks_by_id = {b.id: b for b in Blank.query.filter(
            Blank.id.in_(data['blank_ids'])).all()} if data['blank_ids'] else {}
        for bid in data['blank_ids']:
            b = blanks_by_id[bid]
            db.session.add(ExaminationBlank(
                examination_id=exam.id,
                blank_id=bid,
                quantity=data['blank_qtys'].get(bid, 1),
                price=b.total_price,
                is_insurance=b.is_insurance,
            ))

        db.session.commit()
        flash(f'Barlag №{exam.id} döredilen.', 'success')
        return redirect(url_for('main.examinations_list'))

    defaults = {
        'selected_patient_id': None,
        'selected_patient': None,
        'selected_analysis_qtys': {},
        'selected_direction_ids': set(),
        'selected_tool_qtys': {},
        'selected_blank_qtys': {},
        'doctor_for': {},
    }
    return render_template('main/examinations/create.html', **ctx, **defaults)


# ── Examination detail ────────────────────────────────────────────────────────

@main_bp.route('/examinations/<int:exam_id>')
@examinations_view_required
def examinations_detail(exam_id):
    exam = (
        Examination.query
        .filter_by(id=exam_id)
        .options(
            joinedload(Examination.patient),
            joinedload(Examination.created_by),
            joinedload(Examination.paid_by),
            subqueryload(Examination.exam_analyses).joinedload(ExaminationAnalysis.analysis),
            subqueryload(Examination.exam_directions).joinedload(ExaminationDirection.direction),
            subqueryload(Examination.exam_directions).joinedload(ExaminationDirection.doctor),
            subqueryload(Examination.exam_tools).joinedload(ExaminationAnalysisTool.tool).subqueryload(AnalysisTool.analyses),
            subqueryload(Examination.exam_blanks).joinedload(ExaminationBlank.blank).subqueryload(Blank.analyses),
        )
        .first()
    )
    if exam is None:
        abort(404)

    if current_user.role == 'doctor':
        if not ExaminationDirection.query.filter_by(
            doctor_id=current_user.id, examination_id=exam_id
        ).first():
            flash('Siz üçin bu barlag gadagan.', 'danger')
            return redirect(url_for('main.examinations_list'))

    my_analysis_ids = set()
    if current_user.role == 'analysis_responsible':
        exam_analysis_ids = {ea.analysis_id for ea in exam.exam_analyses}
        my_analysis_ids = {
            a.id for a in Analysis.query.filter(
                Analysis.responsible_id == current_user.id,
                Analysis.id.in_(exam_analysis_ids),
            ).all()
        }
        if not my_analysis_ids:
            flash('Siz üçin bu barlag gadagan.', 'danger')
            return redirect(url_for('main.examinations_list'))

    return render_template('main/examinations/detail.html', exam=exam, my_analysis_ids=my_analysis_ids)


# ── Examination report ────────────────────────────────────────────────────────

@main_bp.route('/examinations/<int:exam_id>/report')
@examinations_view_required
def examinations_report(exam_id):
    exam = (
        Examination.query
        .filter_by(id=exam_id)
        .options(
            joinedload(Examination.patient),
            joinedload(Examination.created_by),
            subqueryload(Examination.exam_analyses).joinedload(ExaminationAnalysis.analysis),
            subqueryload(Examination.exam_directions).joinedload(ExaminationDirection.direction),
            subqueryload(Examination.exam_directions).joinedload(ExaminationDirection.doctor),
            subqueryload(Examination.exam_tools).joinedload(ExaminationAnalysisTool.tool).subqueryload(AnalysisTool.analyses),
            subqueryload(Examination.exam_blanks).joinedload(ExaminationBlank.blank).subqueryload(Blank.analyses),
        )
        .first()
    )
    if exam is None:
        abort(404)

    if current_user.role == 'doctor':
        if not ExaminationDirection.query.filter_by(
            doctor_id=current_user.id, examination_id=exam_id
        ).first():
            abort(403)

    if current_user.role == 'analysis_responsible':
        exam_analysis_ids = {ea.analysis_id for ea in exam.exam_analyses}
        has_own = Analysis.query.filter(
            Analysis.responsible_id == current_user.id,
            Analysis.id.in_(exam_analysis_ids),
        ).first()
        if not has_own:
            abort(403)

    return render_template('main/examinations/report.html', exam=exam)


# ── Edit examination ──────────────────────────────────────────────────────────

@main_bp.route('/examinations/<int:exam_id>/edit', methods=['GET', 'POST'])
@patients_required
def examinations_edit(exam_id):
    exam = (
        Examination.query
        .filter_by(id=exam_id)
        .options(
            joinedload(Examination.patient),
            subqueryload(Examination.exam_analyses),
            subqueryload(Examination.exam_directions),
            subqueryload(Examination.exam_tools),
            subqueryload(Examination.exam_blanks),
        )
        .first()
    )
    if exam is None:
        abort(404)

    if not exam.is_open:
        flash('Diňe açyk barlaglary üýtgedip bolýar.', 'warning')
        return redirect(url_for('main.examinations_list'))

    if exam.is_paid:
        flash('Üýtgetmek mümkin däl-barlag tölenen.', 'warning')
        return redirect(url_for('main.examinations_list'))

    if exam.created_by_id != current_user.id and current_user.role != 'registrar':
        flash('Diňe barlagy döreden üýtgedip bilýär.', 'danger')
        return redirect(url_for('main.examinations_list'))

    ctx = _exam_form_context()

    if request.method == 'POST':
        data, errors = _parse_exam_form()
        if errors:
            for e in errors:
                flash(e, 'danger')
            selected_patient = db.session.get(Patient, data['patient_id']) if data['patient_id'] else None
            return render_template('main/examinations/edit.html', exam=exam, **ctx, **data,
                                   selected_patient=selected_patient)

        patient = db.session.get(Patient, data['patient_id'])
        if exam.patient_id != data['patient_id']:
            exam.patient_has_insurance = bool(patient.insurance_number)
        exam.patient_id = data['patient_id']

        for ea in list(exam.exam_analyses):
            db.session.delete(ea)
        analyses_by_id = {a.id: a for a in Analysis.query.filter(
            Analysis.id.in_(data['analysis_ids'])).all()} if data['analysis_ids'] else {}
        for aid in data['analysis_ids']:
            a = analyses_by_id[aid]
            db.session.add(ExaminationAnalysis(
                examination_id=exam.id,
                analysis_id=aid,
                quantity=data['analysis_qtys'].get(aid, 1),
                price=a.price,
                is_insurance=a.is_insurance,
            ))

        for ed in list(exam.exam_directions):
            db.session.delete(ed)
        directions_by_id = {d.id: d for d in DoctorDirection.query.filter(
            DoctorDirection.id.in_(data['direction_ids'])).all()} if data['direction_ids'] else {}
        for did in data['direction_ids']:
            d = directions_by_id[did]
            db.session.add(ExaminationDirection(
                examination_id=exam.id,
                direction_id=did,
                doctor_id=data['doctor_for'][did],
                price=d.price,
                is_insurance=d.is_insurance,
            ))

        for et in list(exam.exam_tools):
            db.session.delete(et)
        tools_by_id = {t.id: t for t in AnalysisTool.query.filter(
            AnalysisTool.id.in_(data['tool_ids'])).all()} if data['tool_ids'] else {}
        for tid in data['tool_ids']:
            t = tools_by_id[tid]
            db.session.add(ExaminationAnalysisTool(
                examination_id=exam.id,
                tool_id=tid,
                quantity=data['tool_qtys'].get(tid, 1),
                price=t.total_price,
                is_insurance=t.is_insurance,
            ))

        for eb in list(exam.exam_blanks):
            db.session.delete(eb)
        blanks_by_id = {b.id: b for b in Blank.query.filter(
            Blank.id.in_(data['blank_ids'])).all()} if data['blank_ids'] else {}
        for bid in data['blank_ids']:
            b = blanks_by_id[bid]
            db.session.add(ExaminationBlank(
                examination_id=exam.id,
                blank_id=bid,
                quantity=data['blank_qtys'].get(bid, 1),
                price=b.total_price,
                is_insurance=b.is_insurance,
            ))

        db.session.commit()
        flash(f'Barlag №{exam.id} maglumatlary täzelenen.', 'success')
        return redirect(url_for('main.examinations_list'))

    pre = {
        'selected_patient_id': exam.patient_id,
        'selected_patient': exam.patient,
        'selected_analysis_qtys': {ea.analysis_id: ea.quantity for ea in exam.exam_analyses},
        'selected_direction_ids': {ed.direction_id for ed in exam.exam_directions},
        'selected_tool_qtys': {et.tool_id: et.quantity for et in exam.exam_tools},
        'selected_blank_qtys': {eb.blank_id: eb.quantity for eb in exam.exam_blanks},
        'doctor_for': {ed.direction_id: ed.doctor_id for ed in exam.exam_directions},
    }
    return render_template('main/examinations/edit.html', exam=exam, **ctx, **pre)


# ── Close examination ─────────────────────────────────────────────────────────

@main_bp.route('/examinations/<int:exam_id>/close', methods=['POST'])
@patients_required
def examinations_close(exam_id):
    exam = db.session.get(Examination, exam_id)
    if exam is None:
        abort(404)

    if exam.created_by_id != current_user.id and current_user.role != 'registrar':
        flash('Diňe döreden barlagy ýapyp bilýär.', 'danger')
        return redirect(url_for('main.examinations_list'))

    if not exam.is_open:
        flash('Barlag eýýäm ýapylan.', 'warning')
        return redirect(url_for('main.examinations_list'))

    exam.status = Examination.STATUS_CLOSED
    exam.closed_at = datetime.now()
    db.session.commit()
    flash(f'Barlag №{exam.id} ýapylan.', 'success')
    return redirect(url_for('main.examinations_list'))


# ── Toggle insurance discount ─────────────────────────────────────────────────

@main_bp.route('/examinations/<int:exam_id>/toggle-insurance', methods=['POST'])
@role_required('cashier', 'registrar', 'senior_cashier')
def examinations_toggle_insurance(exam_id):
    exam = db.session.get(Examination, exam_id)
    if exam is None:
        abort(404)

    if exam.is_paid:
        flash('Tölenen barlagyň ätiýaçlandyryşyny üýtgedip bolmaýar.', 'warning')
        return redirect(url_for('main.examinations_detail', exam_id=exam_id))

    exam.patient_has_insurance = not exam.patient_has_insurance
    db.session.commit()

    state = 'işjeňleşdirildi' if exam.patient_has_insurance else 'öçürildi'
    flash(f'Ätiýaçlandyryş arzanladyşy {state}.', 'success')
    return redirect(url_for('main.examinations_detail', exam_id=exam_id))


# ── Mark examination as paid ──────────────────────────────────────────────────

@main_bp.route('/examinations/<int:exam_id>/pay', methods=['POST'])
@role_required('cashier', 'senior_cashier')
def examinations_pay(exam_id):
    exam = db.session.get(Examination, exam_id)
    if exam is None:
        abort(404)

    if exam.is_paid:
        flash('Barlag eýýäm tölenen.', 'warning')
        return redirect(url_for('main.examinations_detail', exam_id=exam_id))

    def _pm(field):
        # Default to cash; only an explicit "terminal" switches the method.
        value = request.form.get(field)
        return (ExaminationAnalysis.PAYMENT_TERMINAL
                if value == ExaminationAnalysis.PAYMENT_TERMINAL
                else ExaminationAnalysis.PAYMENT_CASH)

    for ea in exam.exam_analyses:
        ea.payment_method = _pm(f'pm_analysis_{ea.id}')
    for et in exam.exam_tools:
        et.payment_method = _pm(f'pm_tool_{et.id}')
    for eb in exam.exam_blanks:
        eb.payment_method = _pm(f'pm_blank_{eb.id}')
    for ed in exam.exam_directions:
        ed.payment_method = _pm(f'pm_direction_{ed.id}')

    exam.is_paid = True
    exam.paid_at = datetime.now()
    exam.paid_by_id = current_user.id
    db.session.commit()
    flash(f'Barlag №{exam.id} tölenen diýip bellenilen.', 'success')
    return redirect(url_for('main.examinations_detail', exam_id=exam_id))


# ── Mark analysis as submitted ────────────────────────────────────────────────

@main_bp.route('/examinations/<int:exam_id>/analyses/<int:ea_id>/submit', methods=['POST'])
@role_required('analysis_responsible')
def examinations_submit_analysis(exam_id, ea_id):
    ea = (
        ExaminationAnalysis.query
        .filter_by(id=ea_id)
        .options(
            joinedload(ExaminationAnalysis.analysis),
            joinedload(ExaminationAnalysis.examination),
        )
        .first()
    )
    if ea is None or ea.examination_id != exam_id:
        abort(404)

    if ea.analysis.responsible_id != current_user.id:
        flash('Bu analiziň jogapkäri Siz däl.', 'danger')
        return redirect(url_for('main.examinations_detail', exam_id=exam_id))

    if not ea.examination.is_paid:
        flash('Tölenmedik analize bellik goýup bolmaýar.', 'warning')
        return redirect(url_for('main.examinations_detail', exam_id=exam_id))

    if ea.is_submitted:
        flash('Analiz geçilen diýip bellenilen.', 'warning')
        return redirect(url_for('main.examinations_detail', exam_id=exam_id))

    ea.is_submitted = True
    ea.submitted_at = datetime.now()
    ea.submitted_by_id = current_user.id
    db.session.commit()
    flash(f'Analiz «{ea.analysis.name}» geçilen diýip bellenilen.', 'success')
    return redirect(url_for('main.examinations_detail', exam_id=exam_id))


# ── Mark patient visit for direction ─────────────────────────────────────────

@main_bp.route('/examinations/<int:exam_id>/directions/<int:ed_id>/visit', methods=['POST'])
@role_required('doctor')
def examinations_mark_visited(exam_id, ed_id):
    ed = (
        ExaminationDirection.query
        .filter_by(id=ed_id)
        .options(
            joinedload(ExaminationDirection.examination),
            joinedload(ExaminationDirection.direction),
        )
        .first()
    )
    if ed is None or ed.examination_id != exam_id:
        abort(404)

    if ed.doctor_id != current_user.id:
        flash('Siz bu ugur boýunça lukman däl.', 'danger')
        return redirect(url_for('main.examinations_detail', exam_id=exam_id))

    if not ed.examination.is_paid:
        flash('Gelenini belläp bolmaýar-tölenmedik.', 'warning')
        return redirect(url_for('main.examinations_detail', exam_id=exam_id))

    if ed.is_visited:
        flash('Bu ugur boýunça geçilen diýilip bellenilen.', 'warning')
        return redirect(url_for('main.examinations_detail', exam_id=exam_id))

    ed.is_visited = True
    ed.visited_at = datetime.now()
    db.session.commit()
    flash(f'Syrkaw bu ugur geçilen diýip bellenilen «{ed.direction.name}» .', 'success')
    return redirect(url_for('main.examinations_detail', exam_id=exam_id))


# ── Reports (cashier) ─────────────────────────────────────────────────────────

@main_bp.route('/reports')
@reports_required
def reports_index():
    return render_template('main/reports/index.html')


def _new_cat_entry():
    return {'total': Decimal('0'), 'full': Decimal('0'), 'subcategories': OrderedDict()}


def _build_tools_row(exam):
    """Build the detailed tools-report row dict (with category breakdown) for one exam."""
    analyses_total = sum((ea.effective_total for ea in exam.exam_analyses), Decimal('0'))
    directions_total = sum((ed.effective_total for ed in exam.exam_directions), Decimal('0'))
    blanks_total = sum((eb.effective_total for eb in exam.exam_blanks), Decimal('0'))

    analyses_full = sum((ea.snapshot_total for ea in exam.exam_analyses), Decimal('0'))
    directions_full = sum((ed.snapshot_total for ed in exam.exam_directions), Decimal('0'))
    blanks_full = sum((eb.snapshot_total for eb in exam.exam_blanks), Decimal('0'))

    tools_total = Decimal('0')
    tools_full = Decimal('0')
    uncategorized_total = Decimal('0')
    uncategorized_full = Decimal('0')
    categories = OrderedDict()

    for et in exam.exam_tools:
        amount = et.effective_total
        full_amount = et.snapshot_total
        tools_total += amount
        tools_full += full_amount
        tool = et.tool

        if tool.subcategory:
            cat = tool.subcategory.category
            cat_name = cat.name if cat else 'Beleli däl'
            entry = categories.setdefault(cat_name, _new_cat_entry())
            entry['total'] += amount
            entry['full'] += full_amount
            sub_name = tool.subcategory.name
            sub = entry['subcategories'].setdefault(sub_name, {'total': Decimal('0'), 'full': Decimal('0')})
            sub['total'] += amount
            sub['full'] += full_amount
        elif tool.category:
            entry = categories.setdefault(tool.category.name, _new_cat_entry())
            entry['total'] += amount
            entry['full'] += full_amount
        else:
            uncategorized_total += amount
            uncategorized_full += full_amount

    exam_total = analyses_total + directions_total + blanks_total + tools_total
    full_total = analyses_full + directions_full + blanks_full + tools_full
    discount_total = full_total - exam_total

    return {
        'exam': exam,
        'analyses_total': analyses_total,
        'directions_total': directions_total,
        'blanks_total': blanks_total,
        'full_total': full_total,
        'discount_total': discount_total,
        'tools_total': tools_total,
        'exam_total': exam_total,
        'categories': categories,
        'uncategorized_total': uncategorized_total,
        'uncategorized_full': uncategorized_full,
    }


def _reports_grand_totals(query):
    """Aggregate grand totals over the whole filtered query (all pages)."""
    grand = {
        'analyses': Decimal('0'), 'directions': Decimal('0'),
        'blanks': Decimal('0'), 'tools': Decimal('0'), 'total': Decimal('0'),
        'full_total': Decimal('0'), 'discount_total': Decimal('0'),
    }
    exams = (
        query
        .options(
            joinedload(Examination.patient),
            subqueryload(Examination.exam_analyses),
            subqueryload(Examination.exam_directions),
            subqueryload(Examination.exam_blanks),
            subqueryload(Examination.exam_tools),
        )
        .all()
    )
    for exam in exams:
        totals, full = _exam_section_totals(exam)
        grand['analyses'] += totals['analyses']
        grand['directions'] += totals['directions']
        grand['blanks'] += totals['blanks']
        grand['tools'] += totals['tools']
        exam_total = sum(totals.values(), Decimal('0'))
        full_total = sum(full.values(), Decimal('0'))
        grand['total'] += exam_total
        grand['full_total'] += full_total
        grand['discount_total'] += full_total - exam_total
    return grand


def _report_dates_with_today():
    """Return (date_from, date_to) from request args, defaulting both to today
    on the very first load (when no date params are present at all)."""
    if 'from' not in request.args and 'to' not in request.args:
        today = date.today().isoformat()
        return today, today
    return request.args.get('from', '').strip(), request.args.get('to', '').strip()


def _apply_exam_date_filter(query, date_from, date_to):
    """Apply created_at >= from and < to+1day filters; return (query, date_from, date_to)
    with invalid date strings cleared."""
    if date_from:
        try:
            parsed_from = datetime.strptime(date_from, '%Y-%m-%d')
            query = query.filter(Examination.created_at >= parsed_from)
        except ValueError:
            date_from = ''
    if date_to:
        try:
            parsed_to = datetime.strptime(date_to, '%Y-%m-%d') + timedelta(days=1)
            query = query.filter(Examination.created_at < parsed_to)
        except ValueError:
            date_to = ''
    return query, date_from, date_to


def _apply_cashier_scope(query):
    """Restrict a paid-examinations query to the current cashier's own payments.

    A senior cashier sees every cashier's payments, optionally narrowed to one
    cashier via the ?cashier_id query param. Returns (query, cashier_id) where
    cashier_id is the selected filter value (0/None when not filtering)."""
    if current_user.role == 'senior_cashier':
        cashier_id = request.args.get('cashier_id', type=int)
        if cashier_id:
            query = query.filter(Examination.paid_by_id == cashier_id)
        return query, cashier_id
    return query.filter(Examination.paid_by_id == current_user.id), None


def _report_cashiers():
    """Cashiers selectable in the senior-cashier report filter."""
    return (User.query
            .filter(User.role.in_(CASHIER_ROLES))
            .order_by(User.full_name)
            .all())


def _tools_report_query():
    """Build the filtered Examination query for the tools report and return
    (query, date_from, date_to, search, cashier_id). Dates default to today
    on first load."""
    search = request.args.get('q', '').strip()
    date_from, date_to = _report_dates_with_today()

    query = Examination.query.filter(Examination.is_paid == True)
    query, cashier_id = _apply_cashier_scope(query)
    query, date_from, date_to = _apply_exam_date_filter(query, date_from, date_to)

    if search:
        like = f'%{search}%'
        query = query.join(Examination.patient).filter(
            db.or_(
                Patient.full_name.ilike(like),
                Patient.passport_number.ilike(like),
                Patient.insurance_number.ilike(like),
            )
        )

    return query, date_from, date_to, search, cashier_id


def _tools_report_loaded(query):
    return query.options(
        joinedload(Examination.patient),
        joinedload(Examination.created_by),
        joinedload(Examination.paid_by),
        subqueryload(Examination.exam_analyses),
        subqueryload(Examination.exam_directions),
        subqueryload(Examination.exam_blanks),
        subqueryload(Examination.exam_tools)
            .joinedload(ExaminationAnalysisTool.tool)
            .joinedload(AnalysisTool.category),
        subqueryload(Examination.exam_tools)
            .joinedload(ExaminationAnalysisTool.tool)
            .joinedload(AnalysisTool.subcategory)
            .joinedload(AnalysisToolSubcategory.category),
    ).order_by(Examination.created_at.desc())


@main_bp.route('/reports/tools')
@reports_required
def reports_tools():
    query, date_from, date_to, search, cashier_id = _tools_report_query()
    show_cashier = current_user.role == 'senior_cashier'

    if request.args.get('export') == 'xlsx':
        exams = _tools_report_loaded(query).all()
        rows = [_build_tools_row(exam) for exam in exams]
        grand = _reports_grand_totals(query)
        return _tools_report_xlsx(rows, grand, date_from, date_to, show_cashier)

    page = request.args.get('page', 1, type=int)
    pagination = _tools_report_loaded(query).paginate(page=page, per_page=15, error_out=False)
    rows = [_build_tools_row(exam) for exam in pagination.items]
    grand = _reports_grand_totals(query)

    return render_template(
        'main/reports/tools_report.html',
        rows=rows,
        grand=grand,
        pagination=pagination,
        date_from=date_from,
        date_to=date_to,
        search=search,
        show_cashier=show_cashier,
        cashiers=_report_cashiers() if show_cashier else [],
        cashier_id=cashier_id,
    )


def _tools_report_xlsx(rows, grand, date_from, date_to, show_cashier=False):
    """Build an .xlsx workbook of the tools report and return it as a download."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Serişdeler'

    headers = [
        '№', 'Senesi', 'Döreden', 'F.A.A.', 'Doglan ýyly', 'Ýaşy',
        'Raýatlygy', 'Salgysy', 'Pasport №', 'Saglyk ät.№',
        'Analizler', 'Lukmanlar', 'Blanklar', 'Serişdeler',
        'Doly bahasy', 'Ýeňillik', 'Tölenmeli',
    ]

    bold = Font(bold=True)
    muted = Font(color='6C757D')
    header_fill = PatternFill('solid', fgColor='E9ECEF')
    cat_fill = PatternFill('solid', fgColor='F1F3F5')
    right = Alignment(horizontal='right')

    # Breakdown money columns reuse the main "Doly bahasy / Ýeňillik / Tölenmeli" columns.
    BD_FULL, BD_DISC, BD_TOTAL = 15, 16, 17

    def _style_money(row_idx, *cols, bold_cell=False):
        for col in cols:
            c = ws.cell(row=row_idx, column=col)
            c.number_format = '#,##0.00'
            c.alignment = right
            if bold_cell:
                c.font = bold

    def _breakdown_row(name_col, name, full, discount, total):
        row = [''] * len(headers)
        row[name_col - 1] = name
        row[BD_FULL - 1] = float(full)
        row[BD_DISC - 1] = float(discount)
        row[BD_TOTAL - 1] = float(total)
        ws.append(row)
        return ws.max_row

    if show_cashier:
        headers.append('Kassir')

    # Title / period row
    period = ''
    if date_from or date_to:
        period = f'Döwür: {date_from or "…"} — {date_to or "…"}'
    ws.append(['Serişdeler boýunça hasabat'])
    ws['A1'].font = Font(bold=True, size=14)
    if period:
        ws.append([period])
    ws.append([])

    header_row_idx = ws.max_row + 1
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col)
        cell.font = bold
        cell.fill = header_fill

    num_cols = set(range(11, 18))  # money columns
    for r in rows:
        exam = r['exam']
        patient = exam.patient
        ws.append([
            exam.id,
            exam.created_at.strftime('%d.%m.%Y %H:%M'),
            exam.created_by.full_name if exam.created_by else '',
            patient.full_name if patient else '',
            patient.birth_year if patient else '',
            patient.age if patient else '',
            patient.citizenship if patient else '',
            patient.home_address if patient else '',
            (patient.passport_number if patient else '') or '',
            (patient.insurance_number if patient else '') or '',
            float(r['analyses_total']),
            float(r['directions_total']),
            float(r['blanks_total']),
            float(r['tools_total']),
            float(r['full_total']),
            float(r['discount_total']),
            float(r['exam_total']),
        ])
        row_idx = ws.max_row
        for col in num_cols:
            c = ws.cell(row=row_idx, column=col)
            c.number_format = '#,##0.00'
            c.alignment = right
            c.font = bold

        if show_cashier:
            cc = ws.cell(row=row_idx, column=len(headers))
            cc.value = exam.paid_by.full_name if exam.paid_by else ''

        # Category / subcategory breakdown of tools for this examination.
        if r['categories'] or r['uncategorized_total']:
            for cat_name, cat in r['categories'].items():
                cat_disc = cat['full'] - cat['total']
                cidx = _breakdown_row(4, cat_name, cat['full'], cat_disc, cat['total'])
                for col in range(1, len(headers) + 1):
                    ws.cell(row=cidx, column=col).fill = cat_fill
                ws.cell(row=cidx, column=4).font = bold
                _style_money(cidx, BD_FULL, BD_DISC, BD_TOTAL, bold_cell=True)

                for sub_name, sub in cat['subcategories'].items():
                    sub_disc = sub['full'] - sub['total']
                    sidx = _breakdown_row(5, sub_name, sub['full'], sub_disc, sub['total'])
                    ws.cell(row=sidx, column=5).font = muted
                    _style_money(sidx, BD_FULL, BD_DISC, BD_TOTAL)

            if r['uncategorized_total']:
                unc_disc = r['uncategorized_full'] - r['uncategorized_total']
                uidx = _breakdown_row(4, 'Kategoriýasyz serişdeler',
                                      r['uncategorized_full'], unc_disc, r['uncategorized_total'])
                ws.cell(row=uidx, column=4).font = muted
                _style_money(uidx, BD_FULL, BD_DISC, BD_TOTAL)

    # Totals row
    total_row = [
        '', '', '', '', '', '', '', '', '', 'Jemi:',
        float(grand['analyses']), float(grand['directions']),
        float(grand['blanks']), float(grand['tools']),
        float(grand['full_total']), float(grand['discount_total']),
        float(grand['total']),
    ]
    ws.append(total_row)
    total_idx = ws.max_row
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=total_idx, column=col)
        c.font = bold
        if col in num_cols:
            c.number_format = '#,##0.00'
            c.alignment = right

    widths = [6, 16, 20, 24, 10, 6, 14, 28, 14, 14, 11, 11, 11, 11, 12, 11, 12]
    if show_cashier:
        widths = widths + [22]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    fname = 'tools_report'
    if date_from or date_to:
        fname += f'_{date_from or "all"}_{date_to or "all"}'
    fname += '.xlsx'

    return Response(
        buf.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename={fname}'},
    )


def _reports_date_filter():
    """Parse from/to query params and return (query, date_from, date_to)."""
    date_from = request.args.get('from', '').strip()
    date_to = request.args.get('to', '').strip()

    query = Examination.query.filter(Examination.is_paid == True)

    if date_from:
        try:
            parsed_from = datetime.strptime(date_from, '%Y-%m-%d')
            query = query.filter(Examination.created_at >= parsed_from)
        except ValueError:
            date_from = ''
    if date_to:
        try:
            parsed_to = datetime.strptime(date_to, '%Y-%m-%d') + timedelta(days=1)
            query = query.filter(Examination.created_at < parsed_to)
        except ValueError:
            date_to = ''

    return query, date_from, date_to


def _new_grand():
    return {
        'analyses': Decimal('0'), 'directions': Decimal('0'),
        'blanks': Decimal('0'), 'tools': Decimal('0'), 'total': Decimal('0'),
        'full_total': Decimal('0'), 'discount_total': Decimal('0'),
    }


def _exam_section_totals(exam):
    """Return (totals, full) dicts of section sums for an examination."""
    totals = {
        'analyses': sum((ea.effective_total for ea in exam.exam_analyses), Decimal('0')),
        'directions': sum((ed.effective_total for ed in exam.exam_directions), Decimal('0')),
        'blanks': sum((eb.effective_total for eb in exam.exam_blanks), Decimal('0')),
        'tools': sum((et.effective_total for et in exam.exam_tools), Decimal('0')),
    }
    full = {
        'analyses': sum((ea.snapshot_total for ea in exam.exam_analyses), Decimal('0')),
        'directions': sum((ed.snapshot_total for ed in exam.exam_directions), Decimal('0')),
        'blanks': sum((eb.snapshot_total for eb in exam.exam_blanks), Decimal('0')),
        'tools': sum((et.snapshot_total for et in exam.exam_tools), Decimal('0')),
    }
    return totals, full


def _directions_report_query():
    """Build the filtered Examination query for the directions report and return
    (query, date_from, date_to, direction_id, cashier_id). Dates default to today
    on first load; direction_id (when given) keeps only exams that contain that
    doctor-direction."""
    direction_id = request.args.get('direction_id', type=int)
    date_from, date_to = _report_dates_with_today()

    query = Examination.query.filter(Examination.is_paid == True)
    query, cashier_id = _apply_cashier_scope(query)
    query, date_from, date_to = _apply_exam_date_filter(query, date_from, date_to)

    if direction_id:
        query = (
            query
            .join(Examination.exam_directions)
            .filter(ExaminationDirection.direction_id == direction_id)
            .distinct()
        )

    return query, date_from, date_to, direction_id, cashier_id


def _directions_report_loaded(query):
    return query.options(
        joinedload(Examination.patient),
        joinedload(Examination.created_by),
        joinedload(Examination.paid_by),
        subqueryload(Examination.exam_analyses),
        subqueryload(Examination.exam_blanks),
        subqueryload(Examination.exam_tools),
        subqueryload(Examination.exam_directions)
            .joinedload(ExaminationDirection.direction),
        subqueryload(Examination.exam_directions)
            .joinedload(ExaminationDirection.doctor),
    ).order_by(Examination.created_at.desc())


def _build_directions_row(exam):
    """Build the directions-report row dict (with per-direction lines) for one exam."""
    totals, full = _exam_section_totals(exam)

    lines = []
    for ed in exam.exam_directions:
        lines.append({
            'name': ed.direction.name if ed.direction else '—',
            'doctor': ed.doctor.full_name if ed.doctor else '—',
            'visited': ed.is_visited,
            'full': ed.snapshot_total,
            'discount': ed.snapshot_total - ed.effective_total,
            'total': ed.effective_total,
        })

    exam_total = sum(totals.values(), Decimal('0'))
    full_total = sum(full.values(), Decimal('0'))
    discount_total = full_total - exam_total

    return {
        'exam': exam,
        'analyses_total': totals['analyses'],
        'directions_total': totals['directions'],
        'blanks_total': totals['blanks'],
        'tools_total': totals['tools'],
        'full_total': full_total,
        'discount_total': discount_total,
        'exam_total': exam_total,
        'lines': lines,
    }


@main_bp.route('/reports/directions')
@reports_required
def reports_directions():
    query, date_from, date_to, direction_id, cashier_id = _directions_report_query()
    show_cashier = current_user.role == 'senior_cashier'

    if request.args.get('export') == 'xlsx':
        exams = _directions_report_loaded(query).all()
        rows = [_build_directions_row(exam) for exam in exams]
        grand = _reports_grand_totals(query)
        return _directions_report_xlsx(rows, grand, date_from, date_to, show_cashier)

    page = request.args.get('page', 1, type=int)
    pagination = _directions_report_loaded(query).paginate(page=page, per_page=15, error_out=False)
    rows = [_build_directions_row(exam) for exam in pagination.items]
    grand = _reports_grand_totals(query)
    directions = DoctorDirection.query.order_by(DoctorDirection.name).all()

    return render_template(
        'main/reports/directions_report.html',
        rows=rows,
        grand=grand,
        pagination=pagination,
        date_from=date_from,
        date_to=date_to,
        directions=directions,
        direction_id=direction_id,
        show_cashier=show_cashier,
        cashiers=_report_cashiers() if show_cashier else [],
        cashier_id=cashier_id,
    )


def _directions_report_xlsx(rows, grand, date_from, date_to, show_cashier=False):
    """Build an .xlsx workbook of the directions report and return it as a download."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Lukmanlar'

    headers = [
        '№', 'Senesi', 'Döreden', 'F.A.A.', 'Doglan ýyly', 'Ýaşy',
        'Raýatlygy', 'Salgysy', 'Pasport №', 'Saglyk ät.№',
        'Analizler', 'Lukmanlar', 'Blanklar', 'Serişdeler',
        'Doly bahasy', 'Ýeňillik', 'Tölenmeli',
    ]

    bold = Font(bold=True)
    muted = Font(color='6C757D')
    header_fill = PatternFill('solid', fgColor='E9ECEF')
    line_fill = PatternFill('solid', fgColor='F1F3F5')
    right = Alignment(horizontal='right')

    # Breakdown money columns reuse the main "Doly bahasy / Ýeňillik / Tölenmeli" columns.
    BD_FULL, BD_DISC, BD_TOTAL = 15, 16, 17

    def _style_money(row_idx, *cols, bold_cell=False):
        for col in cols:
            c = ws.cell(row=row_idx, column=col)
            c.number_format = '#,##0.00'
            c.alignment = right
            if bold_cell:
                c.font = bold

    if show_cashier:
        headers.append('Kassir')

    # Title / period row
    ws.append(['Lukmanlaryň ugurlary boýunça hasabat'])
    ws['A1'].font = Font(bold=True, size=14)
    if date_from or date_to:
        ws.append([f'Döwür: {date_from or "…"} — {date_to or "…"}'])
    ws.append([])

    header_row_idx = ws.max_row + 1
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col)
        cell.font = bold
        cell.fill = header_fill

    num_cols = set(range(11, 18))  # money columns
    for r in rows:
        exam = r['exam']
        patient = exam.patient
        ws.append([
            exam.id,
            exam.created_at.strftime('%d.%m.%Y %H:%M'),
            exam.created_by.full_name if exam.created_by else '',
            patient.full_name if patient else '',
            patient.birth_year if patient else '',
            patient.age if patient else '',
            patient.citizenship if patient else '',
            patient.home_address if patient else '',
            (patient.passport_number if patient else '') or '',
            (patient.insurance_number if patient else '') or '',
            float(r['analyses_total']),
            float(r['directions_total']),
            float(r['blanks_total']),
            float(r['tools_total']),
            float(r['full_total']),
            float(r['discount_total']),
            float(r['exam_total']),
        ])
        row_idx = ws.max_row
        for col in num_cols:
            c = ws.cell(row=row_idx, column=col)
            c.number_format = '#,##0.00'
            c.alignment = right
            c.font = bold

        if show_cashier:
            cc = ws.cell(row=row_idx, column=len(headers))
            cc.value = exam.paid_by.full_name if exam.paid_by else ''

        # Per-direction breakdown lines for this examination.
        for it in r['lines']:
            line = [''] * len(headers)
            line[3] = it['name']                                  # F.A.A. column → Ugur
            line[4] = it['doctor']                                # Doglan ýyly column → Lukman
            line[5] = 'Baryp gördi' if it['visited'] else 'Garaşylýar'
            line[BD_FULL - 1] = float(it['full'])
            line[BD_DISC - 1] = float(it['discount'])
            line[BD_TOTAL - 1] = float(it['total'])
            ws.append(line)
            lidx = ws.max_row
            for col in range(1, len(headers) + 1):
                ws.cell(row=lidx, column=col).fill = line_fill
            ws.cell(row=lidx, column=4).font = bold
            ws.cell(row=lidx, column=5).font = muted
            ws.cell(row=lidx, column=6).font = muted
            _style_money(lidx, BD_FULL, BD_DISC, BD_TOTAL)

    # Totals row
    total_row = [
        '', '', '', '', '', '', '', '', '', 'Jemi:',
        float(grand['analyses']), float(grand['directions']),
        float(grand['blanks']), float(grand['tools']),
        float(grand['full_total']), float(grand['discount_total']),
        float(grand['total']),
    ]
    ws.append(total_row)
    total_idx = ws.max_row
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=total_idx, column=col)
        c.font = bold
        if col in num_cols:
            c.number_format = '#,##0.00'
            c.alignment = right

    widths = [6, 16, 20, 24, 10, 6, 14, 28, 14, 14, 11, 11, 11, 11, 12, 11, 12]
    if show_cashier:
        widths = widths + [22]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    fname = 'directions_report'
    if date_from or date_to:
        fname += f'_{date_from or "all"}_{date_to or "all"}'
    fname += '.xlsx'

    return Response(
        buf.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename={fname}'},
    )


# ── Doctor-direction categories report (senior cashier only) ──────────────────

_DIRECTION_CATEGORIES_UNCATEGORIZED = 'Kategoriýasyz'


def _direction_categories_report():
    """Aggregate paid doctor-direction income per direction category for the
    selected day (or date range). Returns (rows, totals, date_from, date_to).

    Only paid examinations are considered. Each direction line is split by its
    payment method into cash / terminal buckets (a missing method counts as
    cash, matching the payment default), and its insurance 50% discount amount
    is accumulated separately. Every active category is listed even with no
    income for the period; directions without a category fall into a trailing
    'Kategoriýasyz' bucket that appears only when it has data."""
    date_from, date_to = _report_dates_with_today()

    query = (
        ExaminationDirection.query
        .join(ExaminationDirection.examination)
        .filter(Examination.is_paid == True)
    )
    query, date_from, date_to = _apply_exam_date_filter(query, date_from, date_to)

    lines = (
        query.options(
            joinedload(ExaminationDirection.examination),
            joinedload(ExaminationDirection.direction)
                .joinedload(DoctorDirection.category),
        ).all()
    )

    def _new_bucket(name, is_uncategorized):
        return {
            'name': name,
            'cash': Decimal('0'),
            'terminal': Decimal('0'),
            'discount_cash': Decimal('0'),
            'discount_terminal': Decimal('0'),
            'is_uncategorized': is_uncategorized,
        }

    # Seed every active category so it appears even with no income this period.
    buckets = {
        cat.name: _new_bucket(cat.name, False)
        for cat in DoctorDirectionCategory.query
            .filter(DoctorDirectionCategory.is_active == True)
            .all()
    }

    for ed in lines:
        category = ed.direction.category if ed.direction else None
        name = category.name if category else _DIRECTION_CATEGORIES_UNCATEGORIZED
        bucket = buckets.get(name)
        if bucket is None:
            # An inactive category with income, or the uncategorized bucket.
            bucket = buckets[name] = _new_bucket(name, category is None)

        amount = ed.effective_total
        discount = ed.snapshot_total - ed.effective_total
        if ed.payment_method == ExaminationDirection.PAYMENT_TERMINAL:
            bucket['terminal'] += amount
            bucket['discount_terminal'] += discount
        else:  # cash, or a legacy line with no method — defaults to cash
            bucket['cash'] += amount
            bucket['discount_cash'] += discount

    rows = sorted(
        buckets.values(),
        key=lambda b: (b['is_uncategorized'], b['name'].lower()),
    )
    totals = {
        'cash': sum((b['cash'] for b in rows), Decimal('0')),
        'terminal': sum((b['terminal'] for b in rows), Decimal('0')),
        'discount_cash': sum((b['discount_cash'] for b in rows), Decimal('0')),
        'discount_terminal': sum((b['discount_terminal'] for b in rows), Decimal('0')),
    }
    return rows, totals, date_from, date_to


@main_bp.route('/reports/direction-categories')
@role_required('senior_cashier')
def reports_direction_categories():
    rows, totals, date_from, date_to = _direction_categories_report()

    if request.args.get('export') == 'xlsx':
        return _direction_categories_xlsx(rows, totals, date_from, date_to)

    return render_template(
        'main/reports/direction_categories_report.html',
        rows=rows,
        totals=totals,
        date_from=date_from,
        date_to=date_to,
    )


def _direction_categories_xlsx(rows, totals, date_from, date_to):
    """Build an .xlsx workbook of the direction-categories report."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Kategoriýalar'

    headers = [
        '№',
        'Lukman ugurlarynyň kategoriýasy',
        'Nagt tölegleriň jemi',
        'Terminal tölegleriň jemi',
        '50% ýeňillik (nagt)',
        '50% ýeňillik (terminal)',
        'Nagt tölegleriň umumy jemi',
        'Terminal tölegleriň umumy jemi',
        'Umumy jemi',
    ]

    bold = Font(bold=True)
    header_fill = PatternFill('solid', fgColor='E9ECEF')
    right = Alignment(horizontal='right')
    money_cols = (3, 4, 5, 6, 7, 8, 9)

    ws.append(['Lukman ugurlarynyň kategoriýalary boýunça hasabat'])
    ws['A1'].font = Font(bold=True, size=14)
    if date_from or date_to:
        ws.append([f'Döwür: {date_from or "…"} — {date_to or "…"}'])
    ws.append([])

    header_row_idx = ws.max_row + 1
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col)
        cell.font = bold
        cell.fill = header_fill

    for idx, row in enumerate(rows, start=1):
        ws.append([
            idx,
            row['name'],
            float(row['cash']),
            float(row['terminal']),
            float(row['discount_cash']),
            float(row['discount_terminal']),
            float(row['cash'] + row['discount_cash']),
            float(row['terminal'] + row['discount_terminal']),
            float(row['cash'] + row['discount_cash'] + row['terminal'] + row['discount_terminal']),
        ])
        row_idx = ws.max_row
        for col in money_cols:
            c = ws.cell(row=row_idx, column=col)
            c.number_format = '#,##0.00'
            c.alignment = right

    ws.append([
        '', 'Jemi:',
        float(totals['cash']), float(totals['terminal']),
        float(totals['discount_cash']), float(totals['discount_terminal']),
        float(totals['cash'] + totals['discount_cash']),
        float(totals['terminal'] + totals['discount_terminal']),
        float(totals['cash'] + totals['discount_cash']
              + totals['terminal'] + totals['discount_terminal']),
    ])
    total_idx = ws.max_row
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=total_idx, column=col)
        c.font = bold
        if col in money_cols:
            c.number_format = '#,##0.00'
            c.alignment = right

    widths = [6, 40, 22, 24, 22, 24, 28, 30, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    fname = 'direction_categories_report'
    if date_from or date_to:
        fname += f'_{date_from or "all"}_{date_to or "all"}'
    fname += '.xlsx'

    return Response(
        buf.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename={fname}'},
    )


# ── Doctor / responsible payment report (senior cashier only) ─────────────────

def _doctor_direction_users():
    """Map {direction_id: [User, ...]} of active doctors / analysis-responsibles
    assigned to each doctor-direction, ordered by name. One query."""
    q = (
        db.session.query(user_directions.c.direction_id, User)
        .join(User, User.id == user_directions.c.user_id)
        .filter(User.is_active == True,
                User.role.in_(['doctor', 'analysis_responsible']))
        .order_by(User.full_name)
    )
    mapping = {}
    for direction_id, user in q.all():
        mapping.setdefault(direction_id, []).append(user)
    return mapping


def _doctor_payments_report():
    """Aggregate paid doctor-direction income for one selected direction and one
    selected doctor / responsible over the chosen date range.

    Returns (row, direction, doctor, direction_id, doctor_id, date_from, date_to,
    directions, users_by_direction). `row` is None until both a direction and a
    doctor are picked. Each matching direction line is one reception; its amount
    goes to the cash or terminal bucket by payment method (missing method counts
    as cash), and its insurance 50% discount is accumulated separately."""
    direction_id = request.args.get('direction_id', type=int)
    doctor_id = request.args.get('doctor_id', type=int)
    date_from, date_to = _report_dates_with_today()

    directions = DoctorDirection.query.order_by(DoctorDirection.name).all()
    users_by_direction = _doctor_direction_users()

    row = direction = doctor = None
    if direction_id and doctor_id:
        query = (
            ExaminationDirection.query
            .join(ExaminationDirection.examination)
            .filter(
                Examination.is_paid == True,
                ExaminationDirection.direction_id == direction_id,
                ExaminationDirection.doctor_id == doctor_id,
            )
        )
        query, date_from, date_to = _apply_exam_date_filter(query, date_from, date_to)
        lines = query.all()

        cash = terminal = discount_cash = discount_terminal = Decimal('0')
        for ed in lines:
            amount = ed.effective_total
            discount = ed.snapshot_total - ed.effective_total
            if ed.payment_method == ExaminationDirection.PAYMENT_TERMINAL:
                terminal += amount
                discount_terminal += discount
            else:  # cash, or a legacy line with no method — defaults to cash
                cash += amount
                discount_cash += discount

        row = {
            'count': len(lines),
            'cash': cash,
            'discount_cash': discount_cash,
            'terminal': terminal,
            'discount_terminal': discount_terminal,
        }
        direction = db.session.get(DoctorDirection, direction_id)
        doctor = db.session.get(User, doctor_id)

    return (row, direction, doctor, direction_id, doctor_id,
            date_from, date_to, directions, users_by_direction)


@main_bp.route('/reports/doctor-payments')
@role_required('senior_cashier')
def reports_doctor_payments():
    (row, direction, doctor, direction_id, doctor_id,
     date_from, date_to, directions, users_by_direction) = _doctor_payments_report()

    if request.args.get('export') == 'xlsx' and row is not None:
        return _doctor_payments_xlsx(row, direction, doctor, date_from, date_to)

    # Direction -> eligible users, JSON-serializable for the dependent dropdown.
    users_json = {
        str(did): [[u.id, u.full_name] for u in users]
        for did, users in users_by_direction.items()
    }

    return render_template(
        'main/reports/doctor_payments_report.html',
        row=row,
        direction=direction,
        doctor=doctor,
        direction_id=direction_id,
        doctor_id=doctor_id,
        date_from=date_from,
        date_to=date_to,
        directions=directions,
        users_by_direction=users_by_direction,
        users_json=users_json,
    )


def _doctor_payments_xlsx(row, direction, doctor, date_from, date_to):
    """Build an .xlsx workbook of the doctor / responsible payment report."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Töleg hasabaty'

    headers = [
        'Ugur',
        'Lukman / jogapkär',
        'Kabullaryň sany',
        'Nagt töleg',
        '50% ýeňillik (nagt)',
        'Terminal töleg',
        '50% ýeňillik (terminal)',
        'Nagt tölegleriň umumy jemi',
        'Terminal tölegleriň umumy jemi',
        'Umumy jemi',
    ]

    bold = Font(bold=True)
    header_fill = PatternFill('solid', fgColor='E9ECEF')
    right = Alignment(horizontal='right')
    money_cols = (4, 5, 6, 7, 8, 9, 10)

    ws.append(['Lukman / jogapkär boýunça töleg hasabaty'])
    ws['A1'].font = Font(bold=True, size=14)
    if date_from or date_to:
        ws.append([f'Döwür: {date_from or "…"} — {date_to or "…"}'])
    ws.append([])

    header_row_idx = ws.max_row + 1
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col)
        cell.font = bold
        cell.fill = header_fill

    cash = row['cash']
    discount_cash = row['discount_cash']
    terminal = row['terminal']
    discount_terminal = row['discount_terminal']
    total_cash = cash + discount_cash
    total_terminal = terminal + discount_terminal

    ws.append([
        direction.name if direction else '—',
        doctor.full_name if doctor else '—',
        row['count'],
        float(cash),
        float(discount_cash),
        float(terminal),
        float(discount_terminal),
        float(total_cash),
        float(total_terminal),
        float(total_cash + total_terminal),
    ])
    data_idx = ws.max_row
    ws.cell(row=data_idx, column=3).alignment = right
    for col in money_cols:
        c = ws.cell(row=data_idx, column=col)
        c.number_format = '#,##0.00'
        c.alignment = right

    widths = [30, 26, 16, 16, 20, 16, 22, 28, 30, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    fname = 'doctor_payments_report'
    if date_from or date_to:
        fname += f'_{date_from or "all"}_{date_to or "all"}'
    fname += '.xlsx'

    return Response(
        buf.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename={fname}'},
    )


# ── Analysis-tool payment report (senior cashier only) ────────────────────────

def _tool_payments_report():
    """Aggregate paid analysis-tool income for the selected tools over the
    chosen date range.

    Returns (rows, totals, tool_ids, date_from, date_to, tools). `rows` is None
    until at least one tool is picked; otherwise one bucket per selected tool,
    ordered by tool name. Each matching tool line's amount goes to the cash or
    terminal bucket by payment method (missing method counts as cash), and its
    insurance 50% discount is accumulated separately."""
    tool_ids = list(dict.fromkeys(request.args.getlist('tool_ids', type=int)))
    date_from, date_to = _report_dates_with_today()

    tools = (AnalysisTool.query
             .filter(AnalysisTool.is_active == True)
             .order_by(AnalysisTool.name)
             .all())

    rows = totals = None
    if tool_ids:
        query = (
            ExaminationAnalysisTool.query
            .join(ExaminationAnalysisTool.examination)
            .filter(
                Examination.is_paid == True,
                ExaminationAnalysisTool.tool_id.in_(tool_ids),
            )
            .options(
                joinedload(ExaminationAnalysisTool.examination),
                joinedload(ExaminationAnalysisTool.tool),
            )
        )
        query, date_from, date_to = _apply_exam_date_filter(query, date_from, date_to)
        lines = query.all()

        selected = (AnalysisTool.query
                    .filter(AnalysisTool.id.in_(tool_ids))
                    .order_by(AnalysisTool.name)
                    .all())
        buckets = OrderedDict(
            (t.id, {'name': t.name, 'count': 0,
                    'cash': Decimal('0'), 'discount_cash': Decimal('0'),
                    'terminal': Decimal('0'), 'discount_terminal': Decimal('0')})
            for t in selected
        )
        for et in lines:
            bucket = buckets.get(et.tool_id)
            if bucket is None:
                continue
            amount = et.effective_total
            discount = et.snapshot_total - et.effective_total
            bucket['count'] += et.quantity or 1
            if et.payment_method == ExaminationAnalysisTool.PAYMENT_TERMINAL:
                bucket['terminal'] += amount
                bucket['discount_terminal'] += discount
            else:  # cash, or a legacy line with no method — defaults to cash
                bucket['cash'] += amount
                bucket['discount_cash'] += discount

        rows = list(buckets.values())
        totals = {
            'count': sum(b['count'] for b in rows),
            'cash': sum((b['cash'] for b in rows), Decimal('0')),
            'discount_cash': sum((b['discount_cash'] for b in rows), Decimal('0')),
            'terminal': sum((b['terminal'] for b in rows), Decimal('0')),
            'discount_terminal': sum((b['discount_terminal'] for b in rows), Decimal('0')),
        }

    return rows, totals, tool_ids, date_from, date_to, tools


@main_bp.route('/reports/tool-payments')
@role_required('senior_cashier')
def reports_tool_payments():
    rows, totals, tool_ids, date_from, date_to, tools = _tool_payments_report()

    if request.args.get('export') == 'xlsx' and rows is not None:
        return _tool_payments_xlsx(rows, totals, date_from, date_to)

    return render_template(
        'main/reports/tool_payments_report.html',
        rows=rows,
        totals=totals,
        tool_ids=tool_ids,
        date_from=date_from,
        date_to=date_to,
        tools=tools,
    )


def _tool_payments_xlsx(rows, totals, date_from, date_to):
    """Build an .xlsx workbook of the analysis-tool payment report."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Serişdeler'

    headers = [
        '№',
        'Serişde',
        'Sany',
        'Nagt töleg',
        '50% ýeňillik (nagt)',
        'Terminal töleg',
        '50% ýeňillik (terminal)',
        'Nagt tölegleriň umumy jemi',
        'Terminal tölegleriň umumy jemi',
        'Umumy jemi',
    ]

    bold = Font(bold=True)
    header_fill = PatternFill('solid', fgColor='E9ECEF')
    right = Alignment(horizontal='right')
    money_cols = (4, 5, 6, 7, 8, 9, 10)

    ws.append(['Serişdeler boýunça töleg hasabaty'])
    ws['A1'].font = Font(bold=True, size=14)
    if date_from or date_to:
        ws.append([f'Döwür: {date_from or "…"} — {date_to or "…"}'])
    ws.append([])

    header_row_idx = ws.max_row + 1
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col)
        cell.font = bold
        cell.fill = header_fill

    for idx, row in enumerate(rows, start=1):
        ws.append([
            idx,
            row['name'],
            row['count'],
            float(row['cash']),
            float(row['discount_cash']),
            float(row['terminal']),
            float(row['discount_terminal']),
            float(row['cash'] + row['discount_cash']),
            float(row['terminal'] + row['discount_terminal']),
            float(row['cash'] + row['discount_cash'] + row['terminal'] + row['discount_terminal']),
        ])
        row_idx = ws.max_row
        ws.cell(row=row_idx, column=3).alignment = right
        for col in money_cols:
            c = ws.cell(row=row_idx, column=col)
            c.number_format = '#,##0.00'
            c.alignment = right

    ws.append([
        '', 'Jemi:',
        totals['count'],
        float(totals['cash']), float(totals['discount_cash']),
        float(totals['terminal']), float(totals['discount_terminal']),
        float(totals['cash'] + totals['discount_cash']),
        float(totals['terminal'] + totals['discount_terminal']),
        float(totals['cash'] + totals['discount_cash']
              + totals['terminal'] + totals['discount_terminal']),
    ])
    total_idx = ws.max_row
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=total_idx, column=col)
        c.font = bold
        if col == 3:
            c.alignment = right
        if col in money_cols:
            c.number_format = '#,##0.00'
            c.alignment = right

    widths = [6, 40, 10, 16, 20, 16, 22, 28, 30, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    fname = 'tool_payments_report'
    if date_from or date_to:
        fname += f'_{date_from or "all"}_{date_to or "all"}'
    fname += '.xlsx'

    return Response(
        buf.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename={fname}'},
    )


def _analyses_report_query():
    """Build the filtered Examination query for the analyses report and return
    (query, date_from, date_to, analysis_id, cashier_id). Dates default to today
    on first load; analysis_id (when given) keeps only exams that contain that
    analysis."""
    analysis_id = request.args.get('analysis_id', type=int)
    date_from, date_to = _report_dates_with_today()

    query = Examination.query.filter(Examination.is_paid == True)
    query, cashier_id = _apply_cashier_scope(query)
    query, date_from, date_to = _apply_exam_date_filter(query, date_from, date_to)

    if analysis_id:
        query = (
            query
            .join(Examination.exam_analyses)
            .filter(ExaminationAnalysis.analysis_id == analysis_id)
            .distinct()
        )

    return query, date_from, date_to, analysis_id, cashier_id


def _analyses_report_loaded(query):
    return query.options(
        joinedload(Examination.patient),
        joinedload(Examination.created_by),
        joinedload(Examination.paid_by),
        subqueryload(Examination.exam_directions),
        subqueryload(Examination.exam_blanks),
        subqueryload(Examination.exam_tools),
        subqueryload(Examination.exam_analyses)
            .joinedload(ExaminationAnalysis.analysis)
            .joinedload(Analysis.responsible),
    ).order_by(Examination.created_at.desc())


def _build_analyses_row(exam):
    """Build the analyses-report row dict (with per-analysis lines) for one exam."""
    totals, full = _exam_section_totals(exam)

    lines = []
    for ea in exam.exam_analyses:
        responsible = ea.analysis.responsible if ea.analysis else None
        lines.append({
            'name': ea.analysis.name if ea.analysis else '—',
            'responsible': responsible.full_name if responsible else '—',
            'qty': ea.quantity,
            'submitted': ea.is_submitted,
            'full': ea.snapshot_total,
            'discount': ea.snapshot_total - ea.effective_total,
            'total': ea.effective_total,
        })

    exam_total = sum(totals.values(), Decimal('0'))
    full_total = sum(full.values(), Decimal('0'))
    discount_total = full_total - exam_total

    return {
        'exam': exam,
        'analyses_total': totals['analyses'],
        'directions_total': totals['directions'],
        'blanks_total': totals['blanks'],
        'tools_total': totals['tools'],
        'full_total': full_total,
        'discount_total': discount_total,
        'exam_total': exam_total,
        'lines': lines,
    }


@main_bp.route('/reports/analyses')
@reports_required
def reports_analyses():
    query, date_from, date_to, analysis_id, cashier_id = _analyses_report_query()
    show_cashier = current_user.role == 'senior_cashier'

    if request.args.get('export') == 'xlsx':
        exams = _analyses_report_loaded(query).all()
        rows = [_build_analyses_row(exam) for exam in exams]
        grand = _reports_grand_totals(query)
        return _analyses_report_xlsx(rows, grand, date_from, date_to, show_cashier)

    page = request.args.get('page', 1, type=int)
    pagination = _analyses_report_loaded(query).paginate(page=page, per_page=15, error_out=False)
    rows = [_build_analyses_row(exam) for exam in pagination.items]
    grand = _reports_grand_totals(query)
    analyses = Analysis.query.order_by(Analysis.name).all()

    return render_template(
        'main/reports/analyses_report.html',
        rows=rows,
        grand=grand,
        pagination=pagination,
        date_from=date_from,
        date_to=date_to,
        analyses=analyses,
        analysis_id=analysis_id,
        show_cashier=show_cashier,
        cashiers=_report_cashiers() if show_cashier else [],
        cashier_id=cashier_id,
    )


def _analyses_report_xlsx(rows, grand, date_from, date_to, show_cashier=False):
    """Build an .xlsx workbook of the analyses report and return it as a download."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Analizler'

    headers = [
        '№', 'Senesi', 'Döreden', 'F.A.A.', 'Doglan ýyly', 'Ýaşy',
        'Raýatlygy', 'Salgysy', 'Pasport №', 'Saglyk ät.№',
        'Analizler', 'Lukmanlar', 'Blanklar', 'Serişdeler',
        'Doly bahasy', 'Ýeňillik', 'Tölenmeli',
    ]

    bold = Font(bold=True)
    muted = Font(color='6C757D')
    header_fill = PatternFill('solid', fgColor='E9ECEF')
    line_fill = PatternFill('solid', fgColor='F1F3F5')
    right = Alignment(horizontal='right')
    center = Alignment(horizontal='center')

    # Breakdown money columns reuse the main "Doly bahasy / Ýeňillik / Tölenmeli" columns.
    BD_FULL, BD_DISC, BD_TOTAL = 15, 16, 17

    def _style_money(row_idx, *cols, bold_cell=False):
        for col in cols:
            c = ws.cell(row=row_idx, column=col)
            c.number_format = '#,##0.00'
            c.alignment = right
            if bold_cell:
                c.font = bold

    if show_cashier:
        headers.append('Kassir')

    # Title / period row
    ws.append(['Analizler boýunça hasabat'])
    ws['A1'].font = Font(bold=True, size=14)
    if date_from or date_to:
        ws.append([f'Döwür: {date_from or "…"} — {date_to or "…"}'])
    ws.append([])

    header_row_idx = ws.max_row + 1
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col)
        cell.font = bold
        cell.fill = header_fill

    num_cols = set(range(11, 18))  # money columns
    for r in rows:
        exam = r['exam']
        patient = exam.patient
        ws.append([
            exam.id,
            exam.created_at.strftime('%d.%m.%Y %H:%M'),
            exam.created_by.full_name if exam.created_by else '',
            patient.full_name if patient else '',
            patient.birth_year if patient else '',
            patient.age if patient else '',
            patient.citizenship if patient else '',
            patient.home_address if patient else '',
            (patient.passport_number if patient else '') or '',
            (patient.insurance_number if patient else '') or '',
            float(r['analyses_total']),
            float(r['directions_total']),
            float(r['blanks_total']),
            float(r['tools_total']),
            float(r['full_total']),
            float(r['discount_total']),
            float(r['exam_total']),
        ])
        row_idx = ws.max_row
        for col in num_cols:
            c = ws.cell(row=row_idx, column=col)
            c.number_format = '#,##0.00'
            c.alignment = right
            c.font = bold

        if show_cashier:
            cc = ws.cell(row=row_idx, column=len(headers))
            cc.value = exam.paid_by.full_name if exam.paid_by else ''

        # Per-analysis breakdown lines for this examination.
        for it in r['lines']:
            line = [''] * len(headers)
            line[3] = it['name']                                  # F.A.A. column → Analiz
            line[4] = it['responsible']                           # Doglan ýyly column → Jogapkär
            line[5] = it['qty']                                   # Ýaşy column → Sany
            line[6] = 'Tabşyryldy' if it['submitted'] else 'Garaşylýar'
            line[BD_FULL - 1] = float(it['full'])
            line[BD_DISC - 1] = float(it['discount'])
            line[BD_TOTAL - 1] = float(it['total'])
            ws.append(line)
            lidx = ws.max_row
            for col in range(1, len(headers) + 1):
                ws.cell(row=lidx, column=col).fill = line_fill
            ws.cell(row=lidx, column=4).font = bold
            ws.cell(row=lidx, column=5).font = muted
            ws.cell(row=lidx, column=6).alignment = center
            ws.cell(row=lidx, column=7).font = muted
            _style_money(lidx, BD_FULL, BD_DISC, BD_TOTAL)

    # Totals row
    total_row = [
        '', '', '', '', '', '', '', '', '', 'Jemi:',
        float(grand['analyses']), float(grand['directions']),
        float(grand['blanks']), float(grand['tools']),
        float(grand['full_total']), float(grand['discount_total']),
        float(grand['total']),
    ]
    ws.append(total_row)
    total_idx = ws.max_row
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=total_idx, column=col)
        c.font = bold
        if col in num_cols:
            c.number_format = '#,##0.00'
            c.alignment = right

    widths = [6, 16, 20, 24, 10, 6, 14, 28, 14, 14, 11, 11, 11, 11, 12, 11, 12]
    if show_cashier:
        widths = widths + [22]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    fname = 'analyses_report'
    if date_from or date_to:
        fname += f'_{date_from or "all"}_{date_to or "all"}'
    fname += '.xlsx'

    return Response(
        buf.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename={fname}'},
    )


# ── Earning plans (senior cashier) ────────────────────────────────────────────

PLAN_MONTH_NAMES = {
    1: 'Ýanwar', 2: 'Fewral', 3: 'Mart', 4: 'Aprel', 5: 'Maý', 6: 'Iýun',
    7: 'Iýul', 8: 'Awgust', 9: 'Sentýabr', 10: 'Oktýabr', 11: 'Noýabr', 12: 'Dekabr',
}


def _plan_period():
    """Return (year, month) from request args/form, defaulting to the current month."""
    today = date.today()
    year = request.values.get('year', type=int) or today.year
    month = request.values.get('month', type=int) or today.month
    if month < 1 or month > 12:
        month = today.month
    return year, month


def _month_bounds(year, month):
    """Return (start, next_start) datetimes covering the given calendar month."""
    start = datetime(year, month, 1)
    if month == 12:
        nxt = datetime(year + 1, 1, 1)
    else:
        nxt = datetime(year, month + 1, 1)
    return start, nxt


def _direction_earnings(year, month):
    """Return {user_id: Decimal earned} from paid doctor-directions in the month,
    measured by effective (actually-paid) totals — same semantics as the reports."""
    start, nxt = _month_bounds(year, month)
    lines = (
        ExaminationDirection.query
        .join(ExaminationDirection.examination)
        .filter(
            Examination.is_paid == True,
            Examination.paid_at.isnot(None),
            Examination.paid_at >= start,
            Examination.paid_at < nxt,
        )
        .options(
            joinedload(ExaminationDirection.examination),
            joinedload(ExaminationDirection.direction),
        )
        .all()
    )
    earned = {}
    for ed in lines:
        earned[ed.doctor_id] = earned.get(ed.doctor_id, Decimal('0')) + ed.effective_total
    return earned


def _plan_eligible_users():
    """Doctors / analysis-responsibles who have at least one doctor-direction assigned."""
    return (
        User.query
        .filter(User.role.in_(['doctor', 'analysis_responsible']), User.is_active == True)
        .filter(User.directions.any())
        .order_by(User.full_name)
        .all()
    )


@main_bp.route('/reports/plans')
@role_required('senior_cashier')
def reports_plans():
    year, month = _plan_period()

    eligible = _plan_eligible_users()
    earned = _direction_earnings(year, month)
    plans = {p.user_id: p for p in EarningPlan.query.filter_by(year=year, month=month).all()}

    # Show eligible users plus anyone who already has a plan or earned this month,
    # so no earnings are hidden and past plans stay visible.
    user_ids = set(u.id for u in eligible) | set(plans.keys()) | set(earned.keys())
    users = {u.id: u for u in eligible}
    missing = user_ids - set(users.keys())
    if missing:
        for u in User.query.filter(User.id.in_(missing)).all():
            users[u.id] = u

    rows = []
    plan_total = Decimal('0')
    earned_total = Decimal('0')
    for uid in user_ids:
        user = users.get(uid)
        if user is None:
            continue
        plan = plans.get(uid)
        plan_amount = plan.amount if plan else None
        got = earned.get(uid, Decimal('0'))
        if plan_amount:
            plan_total += plan_amount
        earned_total += got
        pct = None
        if plan_amount and plan_amount > 0:
            pct = float((got / plan_amount) * 100)
        rows.append({
            'user': user,
            'plan': plan_amount,
            'earned': got,
            'diff': got - (plan_amount or Decimal('0')),
            'pct': pct,
        })

    rows.sort(key=lambda r: r['user'].full_name.lower())

    this_year = date.today().year
    years = list(range(this_year - 2, this_year + 2))
    if year not in years:
        years = sorted(set(years) | {year})

    return render_template(
        'main/reports/plans.html',
        rows=rows,
        year=year,
        month=month,
        years=years,
        month_names=PLAN_MONTH_NAMES,
        plan_total=plan_total,
        earned_total=earned_total,
        diff_total=earned_total - plan_total,
    )


@main_bp.route('/reports/plans', methods=['POST'])
@role_required('senior_cashier')
def reports_plans_save():
    year, month = _plan_period()

    for key, raw in request.form.items():
        if not key.startswith('plan_'):
            continue
        try:
            uid = int(key[len('plan_'):])
        except ValueError:
            continue

        user = db.session.get(User, uid)
        if user is None or user.role not in ('doctor', 'analysis_responsible'):
            continue

        plan = EarningPlan.query.filter_by(user_id=uid, year=year, month=month).first()
        value = (raw or '').strip().replace(' ', '').replace(',', '.')

        if not value:
            if plan:
                db.session.delete(plan)
            continue
        try:
            amount = Decimal(value)
        except InvalidOperation:
            continue
        if amount <= 0:
            if plan:
                db.session.delete(plan)
            continue

        if plan:
            plan.amount = amount
        else:
            db.session.add(EarningPlan(user_id=uid, year=year, month=month, amount=amount))

    db.session.commit()
    flash(f'{PLAN_MONTH_NAMES[month]} {year} üçin meýilnama ýatda saklandy.', 'success')
    return redirect(url_for('main.reports_plans', year=year, month=month))
